"""Auditable conversion of professional tide workbook rows into an executable CSV."""
from __future__ import annotations

import csv
import json
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Literal
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import typer
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from clock import shared  # noqa: E402


CONVERTED_FILENAME = "professional_tidal_delta_30s.csv"
DIFF_FILENAME = "professional_tidal_delta_30s.diff.json"
OUTPUT_FILENAMES = (CONVERTED_FILENAME, DIFF_FILENAME)
EXPECTED_HEADERS = {
    "B1": "固体潮",
    "B2": "CAS",
    "C2": "SHA",
    "E1": "海潮",
    "E2": "CAS",
    "F2": "SHA",
    "H2": "CAS-SHA",
}
COMBINED_HEADER_PREFIX = "综合差"
STEP_TOLERANCE_MM = Decimal("1e-9")
DEFAULT_WORKBOOK = Path("clock") / "武汉-上海潮汐结果（0620-0910）-30秒间隔数据.xlsx"
DEFAULT_COMPARE = Path("results") / CONVERTED_FILENAME
XML_NS = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
PKG_NS = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
WB_NS = {
    "a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
app = typer.Typer(add_completion=False, pretty_exceptions_enable=False)


@dataclass(frozen=True, slots=True)
class TideConversionConfig:
    gravity_m_s2: Decimal
    direction: Literal["CAS-minus-SHA"]
    timezone: Literal["UTC"]
    source_column: str = "综合差"
    expected_step_s: int = 30


DEFAULT_CONFIG = TideConversionConfig(
    gravity_m_s2=Decimal("9.794"),
    direction="CAS-minus-SHA",
    timezone="UTC",
)


@dataclass(frozen=True, slots=True)
class ConvertedRow:
    timestamp_utc: str
    total_tidal_delta_m2_s2_surface: Decimal


class ConversionError(ValueError):
    """User-facing validation errors for workbook conversion and comparison."""


def _main_checkout_root() -> Path | None:
    try:
        result = subprocess.run(
            ["git", "-C", str(shared.REPO_ROOT), "rev-parse", "--git-common-dir"],
            check=True,
            capture_output=True,
            text=True,
        )
    except (FileNotFoundError, subprocess.CalledProcessError):
        return None
    common_dir = Path(result.stdout.strip())
    if not common_dir.is_absolute():
        common_dir = (shared.REPO_ROOT / common_dir).resolve()
    else:
        common_dir = common_dir.resolve()
    return common_dir.parent


def _protected_repo_roots() -> tuple[Path, ...]:
    roots = {shared.REPO_ROOT.resolve()}
    main_root = _main_checkout_root()
    if main_root is not None:
        roots.add(main_root.resolve())
    return tuple(sorted(roots))


def _allowed_staging_roots() -> tuple[Path, ...]:
    return tuple(repo_root / "results" / "audit-v1" / "staging" for repo_root in _protected_repo_roots())


def _resolve_cli_path(path: Path) -> Path:
    if path.is_absolute():
        return path.resolve()
    candidate = (Path.cwd() / path).resolve()
    if candidate.exists():
        return candidate
    for root in _protected_repo_roots():
        repo_candidate = (root / path).resolve()
        if repo_candidate.exists():
            return repo_candidate
    return candidate


def _decimal_to_str(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _same_inode(path: Path, other: Path) -> bool:
    try:
        return path.exists() and other.exists() and path.samefile(other)
    except OSError:
        return False


def _validate_output_file(path: Path, *, protected_inputs: tuple[Path, ...]) -> Path:
    if path.is_symlink():
        raise ValueError(f"output filename is a symlink: {path}")
    target = path.parent.resolve(strict=False) / path.name
    resolved_inputs = tuple(input_path.resolve() for input_path in protected_inputs)
    for protected_input in resolved_inputs:
        if target == protected_input or _same_inode(path, protected_input):
            raise ValueError(f"output path collides with protected input: {target}")
    if not any(target.is_relative_to(root) for root in _allowed_staging_roots()):
        raise ValueError(f"output path must stay under results/audit-v1/staging: {target}")
    if path.exists() and (not path.is_file() or path.stat().st_nlink > 1):
        raise ValueError(f"output filename is not an independent regular file: {path}")
    return target


def _write_csv_atomic(path: Path, rows: tuple[ConvertedRow, ...]) -> None:
    with NamedTemporaryFile("w", encoding="utf-8", newline="", delete=False, dir=path.parent) as temp:
        writer = csv.DictWriter(temp, fieldnames=["timestamp_utc", "total_tidal_delta_m2_s2_surface"])
        writer.writeheader()
        for row in rows:
            writer.writerow({
                "timestamp_utc": row.timestamp_utc,
                "total_tidal_delta_m2_s2_surface": _decimal_to_str(row.total_tidal_delta_m2_s2_surface),
            })
        temp_path = Path(temp.name)
    temp_path.replace(path)


def _write_json_atomic(path: Path, payload: dict[str, object]) -> None:
    with NamedTemporaryFile("w", encoding="utf-8", delete=False, dir=path.parent) as temp:
        json.dump(payload, temp, ensure_ascii=False, indent=2)
        temp.write("\n")
        temp_path = Path(temp.name)
    temp_path.replace(path)


def _parse_workbook_timestamp(raw: str) -> str:
    text = raw.strip()
    if text.endswith(".0"):
        text = text[:-2]
    try:
        moment = datetime.strptime(text, "%Y%m%d%H%M%S")
    except ValueError as error:
        raise ConversionError(f"invalid workbook timestamp: {raw}") from error
    return moment.strftime("%Y-%m-%dT%H:%M:%SZ")


def _parse_decimal(raw: str | None, *, label: str) -> Decimal:
    if raw is None or raw.strip() == "":
        raise ConversionError(f"missing numeric value for {label}")
    return Decimal(raw)


def _first_sheet_member(workbook_path: Path) -> str:
    with ZipFile(workbook_path) as archive:
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        rels_xml = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    sheets = workbook_xml.find("a:sheets", WB_NS)
    if sheets is None or not list(sheets):
        raise ConversionError(f"workbook has no sheets: {workbook_path}")
    first_sheet = list(sheets)[0]
    rel_id = first_sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    if rel_id is None:
        raise ConversionError(f"workbook sheet is missing a relationship id: {workbook_path}")
    for rel in rels_xml.findall("r:Relationship", PKG_NS):
        if rel.get("Id") == rel_id:
            target = rel.get("Target")
            if target is None:
                break
            if target.startswith("/"):
                return target.removeprefix("/")
            if target.startswith("xl/"):
                return target
            return f"xl/{target}"
    raise ConversionError(f"could not resolve the first worksheet XML for {workbook_path}")


def _shared_strings(archive: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    values: list[str] = []
    for item in root.findall("a:si", XML_NS):
        values.append("".join(text.text or "" for text in item.iterfind(".//a:t", XML_NS)))
    return values


def _sheet_cells(workbook_path: Path) -> dict[int, dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if not workbook.sheetnames:
            raise ConversionError(f"workbook has no visible sheets: {workbook_path}")
    finally:
        workbook.close()
    sheet_member = _first_sheet_member(workbook_path)
    with ZipFile(workbook_path) as archive:
        shared_strings = _shared_strings(archive)
        root = ET.fromstring(archive.read(sheet_member))
    rows: dict[int, dict[str, str]] = {}
    for row in root.findall(".//a:sheetData/a:row", XML_NS):
        row_index = int(row.attrib["r"])
        values: dict[str, str] = {}
        for cell in row.findall("a:c", XML_NS):
            ref = cell.attrib.get("r")
            if ref is None:
                continue
            column = "".join(ch for ch in ref if ch.isalpha())
            cell_type = cell.attrib.get("t")
            if cell_type == "inlineStr":
                text = "".join(node.text or "" for node in cell.iterfind(".//a:t", XML_NS))
            else:
                value_node = cell.find("a:v", XML_NS)
                if value_node is None or value_node.text is None:
                    continue
                text = value_node.text
                if cell_type == "s":
                    text = shared_strings[int(text)]
            values[column] = text
        if values:
            rows[row_index] = values
    return rows


def _validate_headers(rows: dict[int, dict[str, str]], config: TideConversionConfig) -> None:
    for cell, expected_prefix in EXPECTED_HEADERS.items():
        row_index = int("".join(ch for ch in cell if ch.isdigit()))
        column = "".join(ch for ch in cell if ch.isalpha())
        value = rows.get(row_index, {}).get(column)
        if value is None or not value.startswith(expected_prefix):
            raise ConversionError(f"unexpected workbook header at {cell}: {value!r}")
    combined = rows.get(1, {}).get("J")
    if combined is None or not combined.startswith(config.source_column):
        raise ConversionError(f"unexpected workbook header at J1: {combined!r}")


def convert_height_mm_to_potential(value_mm: Decimal, config: TideConversionConfig) -> Decimal:
    return value_mm * config.gravity_m_s2 / Decimal(1000)


def _rows_from_workbook(workbook_path: Path, config: TideConversionConfig) -> tuple[ConvertedRow, ...]:
    rows = _sheet_cells(workbook_path)
    _validate_headers(rows, config)
    converted: list[ConvertedRow] = []
    previous_stamp: datetime | None = None
    seen_stamps: set[str] = set()
    for row_number in sorted(index for index in rows if index >= 3):
        row = rows[row_number]
        stamp = _parse_workbook_timestamp(row.get("A", ""))
        if stamp in seen_stamps:
            raise ConversionError(f"duplicate UTC timestamp at workbook row {row_number}: {stamp}")
        seen_stamps.add(stamp)
        moment = datetime.strptime(stamp, "%Y-%m-%dT%H:%M:%SZ")
        if previous_stamp is not None:
            delta_s = int((moment - previous_stamp).total_seconds())
            if delta_s != config.expected_step_s:
                raise ConversionError(f"non-{config.expected_step_s}s step at workbook row {row_number}: {delta_s}s")
        previous_stamp = moment
        solid_delta = _parse_decimal(row.get("B"), label=f"B{row_number}") - _parse_decimal(row.get("C"), label=f"C{row_number}")
        ocean_delta = _parse_decimal(row.get("E"), label=f"E{row_number}") - _parse_decimal(row.get("F"), label=f"F{row_number}")
        combined_delta = _parse_decimal(row.get("J"), label=f"J{row_number}")
        if abs((solid_delta + ocean_delta) - combined_delta) > STEP_TOLERANCE_MM:
            raise ConversionError(f"combined difference mismatch at workbook row {row_number}")
        converted.append(
            ConvertedRow(
                timestamp_utc=stamp,
                total_tidal_delta_m2_s2_surface=convert_height_mm_to_potential(combined_delta, config),
            )
        )
    if not converted:
        raise ConversionError(f"no tide rows found in workbook: {workbook_path}")
    return tuple(converted)


def convert_workbook(
    workbook_path: Path,
    output_csv: Path,
    config: TideConversionConfig = DEFAULT_CONFIG,
) -> dict[str, object]:
    workbook_resolved = workbook_path.resolve()
    output_path = _validate_output_file(output_csv, protected_inputs=(workbook_resolved,))
    rows = _rows_from_workbook(workbook_resolved, config)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _write_csv_atomic(output_path, rows)
    return {
        "row_count": len(rows),
        "first_timestamp_utc": rows[0].timestamp_utc,
        "last_timestamp_utc": rows[-1].timestamp_utc,
        "step_seconds": config.expected_step_s,
        "source_column": config.source_column,
    }


def _load_csv_rows(path: Path, *, config: TideConversionConfig) -> tuple[ConvertedRow, ...]:
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    converted: list[ConvertedRow] = []
    previous_stamp: datetime | None = None
    seen_stamps: set[str] = set()
    for row_number, row in enumerate(rows, start=1):
        stamp = row["timestamp_utc"]
        if stamp in seen_stamps:
            raise ConversionError(f"duplicate UTC timestamp at CSV row {row_number}: {stamp}")
        seen_stamps.add(stamp)
        try:
            moment = datetime.strptime(stamp, "%Y-%m-%dT%H:%M:%SZ")
        except ValueError as error:
            raise ConversionError(f"invalid CSV timestamp at row {row_number}: {stamp}") from error
        if previous_stamp is not None:
            delta_s = int((moment - previous_stamp).total_seconds())
            if delta_s != config.expected_step_s:
                raise ConversionError(f"non-{config.expected_step_s}s step at CSV row {row_number}: {delta_s}s")
        previous_stamp = moment
        converted.append(
            ConvertedRow(
                timestamp_utc=stamp,
                total_tidal_delta_m2_s2_surface=Decimal(row["total_tidal_delta_m2_s2_surface"]),
            )
        )
    if not converted:
        raise ConversionError(f"no tide rows found in CSV: {path}")
    return tuple(converted)


def _extra_row_payload(row_number: int, row: ConvertedRow, *, prefix: str) -> dict[str, object]:
    return {
        "row_number": row_number,
        f"{prefix}_timestamp_utc": row.timestamp_utc,
        f"{prefix}_value": _decimal_to_str(row.total_tidal_delta_m2_s2_surface),
    }


def compare_tide_csv(
    converted_csv: Path,
    reference_csv: Path,
    config: TideConversionConfig = DEFAULT_CONFIG,
) -> dict[str, object]:
    converted_rows = _load_csv_rows(converted_csv, config=config)
    reference_rows = _load_csv_rows(reference_csv, config=config)
    common_prefix_row_count = min(len(converted_rows), len(reference_rows))
    max_abs_time_diff_s = 0
    max_abs_value_diff = Decimal(0)
    first_inconsistent_row: dict[str, object] | None = None
    for row_number in range(1, common_prefix_row_count + 1):
        converted = converted_rows[row_number - 1]
        reference = reference_rows[row_number - 1]
        converted_moment = datetime.strptime(converted.timestamp_utc, "%Y-%m-%dT%H:%M:%SZ")
        reference_moment = datetime.strptime(reference.timestamp_utc, "%Y-%m-%dT%H:%M:%SZ")
        time_diff_s = int((converted_moment - reference_moment).total_seconds())
        value_diff = converted.total_tidal_delta_m2_s2_surface - reference.total_tidal_delta_m2_s2_surface
        max_abs_time_diff_s = max(max_abs_time_diff_s, abs(time_diff_s))
        max_abs_value_diff = max(max_abs_value_diff, abs(value_diff))
        if first_inconsistent_row is None and (time_diff_s != 0 or value_diff != 0):
            first_inconsistent_row = {
                "row_number": row_number,
                "converted_timestamp_utc": converted.timestamp_utc,
                "reference_timestamp_utc": reference.timestamp_utc,
                "converted_value": _decimal_to_str(converted.total_tidal_delta_m2_s2_surface),
                "reference_value": _decimal_to_str(reference.total_tidal_delta_m2_s2_surface),
                "time_diff_s": abs(time_diff_s),
                "value_diff": _decimal_to_str(value_diff),
            }
    first_extra_converted_row = None
    if len(converted_rows) > common_prefix_row_count:
        first_extra_converted_row = _extra_row_payload(
            common_prefix_row_count + 1,
            converted_rows[common_prefix_row_count],
            prefix="converted",
        )
    first_extra_reference_row = None
    if len(reference_rows) > common_prefix_row_count:
        first_extra_reference_row = _extra_row_payload(
            common_prefix_row_count + 1,
            reference_rows[common_prefix_row_count],
            prefix="reference",
        )
    return {
        "converted_row_count": len(converted_rows),
        "reference_row_count": len(reference_rows),
        "common_prefix_row_count": common_prefix_row_count,
        "first_timestamp_utc": converted_rows[0].timestamp_utc,
        "last_timestamp_utc": converted_rows[-1].timestamp_utc,
        "max_abs_time_diff_s": max_abs_time_diff_s,
        "max_abs_value_diff": _decimal_to_str(max_abs_value_diff),
        "first_inconsistent_row": first_inconsistent_row,
        "first_extra_converted_row": first_extra_converted_row,
        "first_extra_reference_row": first_extra_reference_row,
    }


def _comparison_has_mismatch(summary: dict[str, object]) -> bool:
    return any(
        summary[key] is not None
        for key in ("first_inconsistent_row", "first_extra_converted_row", "first_extra_reference_row")
    )


@app.command()
def main(
    input: Path = typer.Option(DEFAULT_WORKBOOK, "--input", exists=False, dir_okay=False),
    compare: Path = typer.Option(DEFAULT_COMPARE, "--compare", exists=False, dir_okay=False),
    output_dir: Path = typer.Option(..., "--output-dir", exists=False, file_okay=False, dir_okay=True),
    tolerate_mismatch: bool = typer.Option(False, "--tolerate-mismatch", help="Record a workbook/CSV mismatch in the diff report and exit 0 (audit mode)."),
) -> None:
    try:
        input_path = _resolve_cli_path(input)
        compare_path = _resolve_cli_path(compare)
        output_root = output_dir.resolve()
        converted_path = output_root / CONVERTED_FILENAME
        diff_path = output_root / DIFF_FILENAME
        conversion_summary = convert_workbook(input_path, converted_path)
        comparison_summary = compare_tide_csv(converted_path, compare_path)
        _validate_output_file(diff_path, protected_inputs=(input_path.resolve(), compare_path.resolve(), converted_path.resolve()))
        diff_path.parent.mkdir(parents=True, exist_ok=True)
        report = {
            "config": {
                "gravity_m_s2": _decimal_to_str(DEFAULT_CONFIG.gravity_m_s2),
                "direction": DEFAULT_CONFIG.direction,
                "timezone": DEFAULT_CONFIG.timezone,
                "source_column": DEFAULT_CONFIG.source_column,
                "expected_step_s": DEFAULT_CONFIG.expected_step_s,
            },
            "input_workbook": str(input_path),
            "reference_csv": str(compare_path),
            "conversion_summary": conversion_summary,
            "comparison_summary": comparison_summary,
        }
        _write_json_atomic(diff_path, report)
        if _comparison_has_mismatch(comparison_summary):
            if tolerate_mismatch:
                typer.echo(
                    "Warning: converted workbook does not match the tracked CSV; "
                    "mismatch recorded in the diff report",
                    err=True,
                )
            else:
                raise ConversionError(
                    "converted workbook does not match the tracked CSV; see diff report for the first inconsistent or extra row"
                )
    except (ConversionError, OSError, ValueError) as error:
        typer.echo(f"Error: {error}", err=True)
        raise typer.Exit(code=1) from error
    typer.echo(f"Wrote converted tide CSV and diff report to {output_root}")


if __name__ == "__main__":
    app()
